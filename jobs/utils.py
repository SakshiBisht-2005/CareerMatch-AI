
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from .pdf_processor import extract_pdf_text

def extract_resume_text(file_path):
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        result = extract_pdf_text(path)
        return result["text"]

    if extension == ".docx":
        document = Document(path)

        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text for cell in row.cells
                    )
                )

        return "\n".join(parts).strip()

    raise ValueError(
        "Only PDF and DOCX resumes are supported."
    )


def detect_skills(text):
    skill_list = [
        "Python",
        "Django",
        "Java",
        "JavaScript",
        "HTML",
        "CSS",
        "SQL",
        "MySQL",
        "PostgreSQL",
        "React",
        "Node.js",
        "Machine Learning",
        "Deep Learning",
        "Data Analysis",
        "Pandas",
        "NumPy",
        "Scikit-learn",
        "TensorFlow",
        "PyTorch",
        "Git",
        "GitHub",
        "REST API",
        "Docker",
        "AWS",
        "Excel",
    ]

    text_lower = text.lower()
    detected = []

    for skill in skill_list:
        if skill.lower() in text_lower:
            detected.append(skill)

    return detected
